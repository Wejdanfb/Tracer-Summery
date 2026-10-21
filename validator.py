"""
validator.py
Re-opens the generated Summary with openpyxl and checks it against the
extracted tracer data and the original template formatting.
"""
import openpyxl

from excel_utils import col_letter
from summary_writer import LAYOUT_PER_UNIT, NO_FINDINGS_TEXT, pick_template_sheet, find_template_layout


def _style_sig(cell):
    f, b, fl, a = cell.font, cell.border, cell.fill, cell.alignment
    return (f.name, f.sz, f.b, f.i, str(f.color.rgb) if f.color is not None else None,
            b.left.style, b.right.style, b.top.style, b.bottom.style,
            fl.fill_type, str(fl.fgColor.rgb) if fl.fgColor is not None else None,
            a.horizontal)


def validate_output(output_path, template_path, results, info):
    checks = []

    def check(ok, msg):
        checks.append((bool(ok), msg))

    try:
        wb = openpyxl.load_workbook(output_path)
        wb_vals = openpyxl.load_workbook(output_path, data_only=True)
    except Exception as e:
        return [(False, "Output workbook could not be opened: %s" % e)]
    check(True, "Workbook opens")
    tpl = openpyxl.load_workbook(template_path)
    ws = wb[info["sheet"]]
    tws = pick_template_sheet(tpl)
    cols = info["cols"]
    first, last = info["first_data"], info["last_data"]
    layout = info["layout"]

    check(wb.sheetnames == tpl.sheetnames, "Sheet names preserved %s" % wb.sheetnames)

    # ---- expected rows ----
    expected = []
    for res in results:
        texts = [it.summary_text for it in res.findings]
        if layout == LAYOUT_PER_UNIT:
            expected.append((res.unit, res.final_score, "\n".join(texts) if texts else NO_FINDINGS_TEXT))
        else:
            for t in (texts or [NO_FINDINGS_TEXT]):
                expected.append((res.unit, res.final_score, t))
    actual = []
    for r in range(first, last + 1):
        actual.append((ws.cell(row=r, column=cols["unit"]).value,
                       ws.cell(row=r, column=cols["score"]).value,
                       ws.cell(row=r, column=cols["findings"]).value))
    check(len(actual) == len(expected), "Row count: %d written / %d expected" % (len(actual), len(expected)))
    if layout == LAYOUT_PER_UNIT:
        check(len(actual) == len(results),
              "One row per unit/ward: %d rows for %d tracer files" % (len(actual), len(results)))
        per_line_ok = all(
            (a[2] or "") == "\n".join([it.summary_text for it in res.findings] or [NO_FINDINGS_TEXT])
            for a, res in zip(actual, results)) if len(actual) == len(results) else False
        check(per_line_ok, "All findings of each unit in ONE cell, one finding per line")
        wrap_ok = all(ws.cell(row=r, column=cols["findings"]).alignment.wrap_text for r in range(first, last + 1))
        check(wrap_ok, "Wrap Text enabled on Significant Findings cells")

    # units
    units_out = [ws.cell(row=r, column=cols["unit"]).value for r in info["unit_first_rows"]]
    units_exp = [res.unit for res in results]
    check(units_out == units_exp, "Units (%d) and order correct: %s" % (len(units_exp), ", ".join(map(str, units_exp))))

    # scores
    bad_scores = [(i + first, a[1], e[1]) for i, (a, e) in enumerate(zip(actual, expected))
                  if a[1] is None or e[1] is None or abs(float(a[1]) - float(e[1])) > 1e-9]
    check(not bad_scores, "Tracer scores correct" + ("" if not bad_scores else " - mismatches at rows %s" % bad_scores[:5]))

    # findings text
    bad_text = [i + first for i, (a, e) in enumerate(zip(actual, expected)) if a[2] != e[2]]
    check(not bad_text, "Significant Findings text matches source exactly"
          + ("" if not bad_text else " - mismatches at rows %s" % bad_text[:10]))

    # inclusion / exclusion rules (checked explicitly per unit)
    out_by_unit = {}
    for u, _, t in actual:
        out_by_unit.setdefault(u, []).append(t or "")
    n0 = n1 = n_notes = n_desc = 0
    missing, leaked = [], []
    for res in results:
        blob = "\n".join(out_by_unit.get(res.unit, []))
        for it in res.items:
            if it.score in (0, 1):
                n0 += it.score == 0
                n1 += it.score == 1
                if it.has_notes:
                    n_notes += 1
                else:
                    n_desc += 1
                if it.summary_text not in blob:
                    missing.append("%s #%d" % (res.unit, it.number))
            elif it.score in (2, "NA"):
                # an excluded item must not appear (unless identical text is also an included finding)
                included_texts = {x.finding_text for x in res.findings}
                for txt in filter(None, [it.notes if it.has_notes else None, it.description]):
                    if txt in blob and txt not in included_texts and not any(txt in x for x in included_texts):
                        leaked.append("%s #%d" % (res.unit, it.number))
    check(not missing, "All score 0 (%d) and score 1 (%d) findings included" % (n0, n1)
          + ("" if not missing else " - MISSING: %s" % missing[:10]))
    # STD No. present in front of every 0/1 finding
    no_std, n_std = [], 0
    lines = []
    for t in (a[2] for a in actual):
        lines.extend((t or "").split("\n") if layout == LAYOUT_PER_UNIT else [t or ""])
    for res in results:
        for it in res.findings:
            if it.std_label and ("%s — %s" % (it.std_label, it.finding_text)) in "\n".join(out_by_unit.get(res.unit, [])):
                n_std += 1
            else:
                no_std.append("%s #%d" % (res.unit, it.number))
    check(not no_std, "STD No. shown before every finding (%d of %d)" % (n_std, n0 + n1)
          + ("" if not no_std else " - WITHOUT STD: %s" % no_std[:10]))
    check(not leaked, "Score 2 and NA items excluded" + ("" if not leaked else " - FOUND: %s" % leaked[:10]))
    check(not bad_text, "Notes transferred verbatim (%d) / empty Notes replaced by checklist description (%d)"
          % (n_notes, n_desc))

    # Action plan / evidence blank
    for key, label in (("action", "Action Plan"), ("evidence", "Evidence")):
        if key in cols:
            filled = [r for r in range(first, last + 1) if ws.cell(row=r, column=cols[key]).value not in (None, "")]
            check(not filled, "%s column blank" % label + ("" if not filled else " - filled at rows %s" % filled))

    # SN sequential
    sns = [ws.cell(row=r, column=cols["sn"]).value for r in range(first, last + 1)]
    check(sns == list(range(1, len(sns) + 1)), "SN sequential 1..%d" % len(sns))

    # total row
    if info.get("total_row"):
        tv = ws.cell(row=info["total_row"], column=cols["score"]).value
        check(isinstance(tv, (str, float, int)) and tv not in (None, ""), "TOTAL SCORE row present (row %d): %s" % (info["total_row"], tv))

    # ---- formatting ----
    th, tcols, tfirst_col, tlast_col, ttotal, _ = find_template_layout(tws)
    hdr_ok = all(ws.cell(row=th, column=c).value == tws.cell(row=th, column=c).value and
                 _style_sig(ws.cell(row=th, column=c)) == _style_sig(tws.cell(row=th, column=c))
                 for c in range(tfirst_col, tlast_col + 1))
    check(hdr_ok, "Header row text and style unchanged")

    skip = info.get("month_cell")
    top_ok = all(ws.cell(row=r, column=c).value == tws.cell(row=r, column=c).value
                 for r in range(1, th) for c in range(1, tws.max_column + 1)
                 if ws.cell(row=r, column=c).coordinate != skip)
    check(top_ok, "Title block / NPI / Month rows unchanged" + (" (except MONTH\\YEAR %s, set by user)" % skip if skip else ""))

    tmerged = {str(m) for m in tws.merged_cells.ranges if m.max_row < th + 1}
    omerged = {str(m) for m in ws.merged_cells.ranges if m.max_row < th + 1}
    check(tmerged == omerged, "Merged cells preserved: %s" % ", ".join(sorted(tmerged)))

    cw_ok = all((ws.column_dimensions[k].width == v.width) for k, v in tws.column_dimensions.items())
    check(cw_ok, "Column widths preserved")
    rh_ok = all(ws.row_dimensions[r].height == tws.row_dimensions[r].height for r in range(1, th + 1))
    check(rh_ok, "Header/title row heights preserved")
    check(len(ws._images) == len(tws._images), "Logo/images preserved (%d)" % len(ws._images))
    check(ws.page_setup.orientation == tws.page_setup.orientation and
          ws.page_setup.paperSize == tws.page_setup.paperSize and
          ws.page_margins.left == tws.page_margins.left, "Page setup preserved")

    proto_sig = [_style_sig(tws.cell(row=th + 1, column=c)) for c in range(tfirst_col, tlast_col + 1)]
    style_bad = []
    for r in range(first, last + 1):
        sig = [_style_sig(ws.cell(row=r, column=c)) for c in range(tfirst_col, tlast_col + 1)]
        # findings alignment is forced to wrap/left; compare everything except alignment for that column
        fi = cols["findings"] - tfirst_col
        sig[fi] = sig[fi][:-1]
        exp = list(proto_sig)
        exp[fi] = exp[fi][:-1]
        if sig != exp:
            style_bad.append(r)
    check(not style_bad, "Data rows use template fonts/borders/fills" + ("" if not style_bad else " - rows %s" % style_bad[:5]))
    pct_ok = all("%" in ws.cell(row=r, column=cols["score"]).number_format for r in range(first, last + 1))
    check(pct_ok, "Tracer Score cells formatted as percentage")
    return checks
