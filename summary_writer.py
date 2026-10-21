"""
summary_writer.py
Fills the Tracer Summary TEMPLATE with the extracted findings and saves a NEW workbook.

Template structure (detected dynamically, observed in TRACER.SUMMERY.MCH.2026):
  * Logo + title block (merged A1:F1, A2:F5, A6:F6)  -> untouched
  * Row 7: NPI NAME / MONTH\\YEAR                   -> untouched (optional month override)
  * Header row: SN | Unit/ Ward | Tracer Score (%) | Significant Findings |
                Action Plan Submitted | Evidence of Implementation Submitted
  * Example data rows                                -> cleared and replaced
  * "TOTAL SCORE :" row with =AVERAGE(...)           -> moved below the new data
  * Helper lists (Yes/No in Z1:Z2, month list in T14:T25) -> untouched
"""
import re
from copy import copy
from dataclasses import dataclass
from typing import List, Optional

import openpyxl
from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.cell_range import CellRange

from excel_utils import norm_text, copy_cell_style, estimate_row_height, col_letter

LAYOUT_PER_FINDING = "per_finding"   # one row per finding (unit + score repeated)
LAYOUT_PER_UNIT = "per_unit"         # one row per unit, findings listed in one cell

NO_FINDINGS_TEXT = "/"  # used when a unit has no item scored 0 or 1 (template convention)

TEMPLATE_HEADERS = {
    "sn": [r"^s\.?\s?n\.?$", r"^serial", r"^no\.?$"],
    "unit": [r"unit", r"ward", r"department"],
    "score": [r"score"],
    "findings": [r"significant", r"finding"],
    "action": [r"action\s*plan"],
    "evidence": [r"evidence"],
}


@dataclass
class OutRow:
    sn: int
    unit: str
    score: float
    finding: str
    first_of_unit: bool
    source_file: str


class TemplateError(Exception):
    pass


# ---------------------------------------------------------------------------
def find_template_layout(ws, max_rows=60):
    for r in range(1, min(ws.max_row, max_rows) + 1):
        cols = {}
        for c in range(1, min(ws.max_column, 30) + 1):
            t = norm_text(ws.cell(row=r, column=c).value)
            if not t:
                continue
            for key in ("sn", "action", "evidence", "findings", "score", "unit"):
                if key not in cols and any(re.search(p, t) for p in TEMPLATE_HEADERS[key]):
                    cols[key] = c
                    break
        if {"sn", "unit", "score", "findings"} <= set(cols):
            header_row = r
            break
    else:
        raise TemplateError("Could not find the Summary header row (SN / Unit / Tracer Score / Significant Findings).")

    first_col = min(cols.values())
    last_col = max(cols.values())
    total_row = None
    for r in range(header_row + 1, ws.max_row + 1):
        for c in range(first_col, last_col + 1):
            if re.search(r"total", norm_text(ws.cell(row=r, column=c).value)):
                total_row = r
                break
        if total_row:
            break
    if total_row is None:
        # no total row: data block ends at last row with content
        last = header_row
        for r in range(header_row + 1, ws.max_row + 1):
            if any(ws.cell(row=r, column=c).value not in (None, "") for c in range(first_col, last_col + 1)):
                last = r
        old_end = last
    else:
        old_end = total_row
    return header_row, cols, first_col, last_col, total_row, old_end


def pick_template_sheet(wb):
    for ws in wb.worksheets:
        try:
            find_template_layout(ws)
            return ws
        except TemplateError:
            continue
    raise TemplateError("No worksheet in the template contains the Summary header "
                        "(SN / Unit/Ward / Tracer Score / Significant Findings).")


# ---------------------------------------------------------------------------
def build_rows(results, layout):
    rows: List[OutRow] = []
    sn = 0
    for res in results:
        texts = [it.summary_text for it in res.findings]
        if layout == LAYOUT_PER_UNIT:
            sn += 1
            txt = "\n".join(texts) if texts else NO_FINDINGS_TEXT
            rows.append(OutRow(sn, res.unit, res.final_score, txt, True, res.file_name))
        else:
            if not texts:
                texts = [NO_FINDINGS_TEXT]
            for i, t in enumerate(texts):
                sn += 1
                rows.append(OutRow(sn, res.unit, res.final_score, t, i == 0, res.file_name))
    return rows


def _snapshot_row(ws, row, first_col, last_col):
    cells = {}
    for c in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cells[c] = (copy(cell._style), cell.value, cell.number_format)
    h = ws.row_dimensions[row].height if row in ws.row_dimensions else None
    return cells, h


def _apply_snapshot_style(ws, row, snap, first_col, last_col):
    cells, _ = snap
    for c in range(first_col, last_col + 1):
        ws.cell(row=row, column=c)._style = copy(cells[c][0])


def write_summary(template_path, results, output_path, layout=LAYOUT_PER_UNIT, month_value=None, log=print):
    wb = openpyxl.load_workbook(template_path)
    ws = pick_template_sheet(wb)
    header_row, cols, first_col, last_col, total_row, old_end = find_template_layout(ws)
    first_data = header_row + 1
    log("Template sheet '%s': header row %d, data rows %d-%d, total row %s"
        % (ws.title, header_row, first_data, (total_row or old_end + 1) - 1, total_row or "none"))

    # ---- snapshots of prototype rows (styles/heights) ----
    proto = _snapshot_row(ws, first_data, first_col, last_col)
    total_snap = _snapshot_row(ws, total_row, first_col, last_col) if total_row else None
    blank_src = (total_row or old_end) + 1
    blank_snap = _snapshot_row(ws, blank_src, first_col, last_col)
    default_h = ws.sheet_format.defaultRowHeight or 15

    # ---- remove merges inside the data block (template has none, but be safe) ----
    for rng in list(ws.merged_cells.ranges):
        if rng.max_row >= first_data and rng.min_row <= old_end and rng.max_col >= first_col and rng.min_col <= last_col:
            ws.unmerge_cells(str(rng))

    # ---- clear old example rows ----
    for r in range(first_data, old_end + 1):
        for c in range(first_col, last_col + 1):
            ws.cell(row=r, column=c).value = None

    rows = build_rows(results, layout)
    widths = {c: (ws.column_dimensions[col_letter(c)].width or 10) for c in range(first_col, last_col + 1)}

    # ---- write data rows ----
    r = first_data
    unit_first_rows = []
    for orow in rows:
        _apply_snapshot_style(ws, r, proto, first_col, last_col)
        ws.cell(row=r, column=cols["sn"]).value = orow.sn
        ws.cell(row=r, column=cols["unit"]).value = orow.unit
        sc = ws.cell(row=r, column=cols["score"])
        sc.value = orow.score
        if "%" not in (sc.number_format or ""):
            sc.number_format = "0.00%"
        fc = ws.cell(row=r, column=cols["findings"])
        fc.value = orow.finding
        al = copy(fc.alignment)
        fc.alignment = Alignment(horizontal=al.horizontal or "left", vertical=al.vertical or "center",
                                 wrap_text=True, indent=al.indent, text_rotation=al.text_rotation)
        for key in ("action", "evidence"):
            if key in cols:
                ws.cell(row=r, column=cols[key]).value = None  # ALWAYS blank
        needed = estimate_row_height(
            [(orow.finding, widths[cols["findings"]]), (orow.unit, widths[cols["unit"]])],
            min_height=max(30.0, proto[1] or 30.0) if layout == LAYOUT_PER_UNIT else 30.0,
            max_height=100000)
        # Excel's maximum row height is 409 points (about 27 lines of text)
        ws.row_dimensions[r].height = min(needed, 409.0)
        if needed > 409.0:
            log("WARNING: %s has many findings; Excel's maximum row height (409 pt) is reached - "
                "some lines may need the cell to be selected/zoomed to read." % orow.unit)
        if orow.first_of_unit:
            unit_first_rows.append(r)
        r += 1
    last_data = r - 1

    # ---- total row ----
    new_total = None
    if total_snap:
        new_total = r
        _apply_snapshot_style(ws, new_total, total_snap, first_col, last_col)
        for c in range(first_col, last_col + 1):
            v = total_snap[0][c][1]
            ws.cell(row=new_total, column=c).value = None if (isinstance(v, str) and v.startswith("=")) else v
        sc_col = col_letter(cols["score"])
        if unit_first_rows:
            if layout == LAYOUT_PER_UNIT or len(unit_first_rows) == len(rows):
                formula = "=AVERAGE(%s%d:%s%d)" % (sc_col, first_data, sc_col, last_data)
            elif len(unit_first_rows) <= 250:
                # one score per UNIT (not per finding) so units with many findings are not over-weighted
                formula = "=AVERAGE(%s)" % ",".join("%s%d" % (sc_col, x) for x in unit_first_rows)
            else:
                formula = None
            tc = ws.cell(row=new_total, column=cols["score"])
            if formula:
                tc.value = formula
            else:
                vals = [rw.score for rw in rows if rw.first_of_unit]
                tc.value = sum(vals) / len(vals)
        if total_snap[1]:
            ws.row_dimensions[new_total].height = total_snap[1]
        r = new_total + 1

    # ---- rows that previously held template data but are now unused ----
    for rr in range(r, old_end + 1):
        _apply_snapshot_style(ws, rr, blank_snap, first_col, last_col)
        for c in range(first_col, last_col + 1):
            ws.cell(row=rr, column=c).value = None
        ws.row_dimensions[rr].height = blank_snap[1] or default_h

    # ---- data validation (Yes/No dropdown) for Action Plan / Evidence columns ----
    _fix_validations(ws, cols, first_data, old_end, last_data)

    # ---- print area (if the template has one) ----
    end_row = new_total or last_data
    if ws.print_area:
        m = re.findall(r"\$?([A-Z]+)\$?(\d+)", ws.print_area.split("!")[-1])
        if len(m) >= 2:
            ws.print_area = "%s%s:%s%d" % (m[0][0], m[0][1], m[-1][0], max(int(m[-1][1]), end_row))

    # ---- optional MONTH\YEAR override ----
    month_cell = _set_month(ws, header_row, month_value, log) if month_value is not None else None

    wb.save(output_path)
    return {
        "sheet": ws.title, "header_row": header_row, "cols": cols, "first_data": first_data,
        "last_data": last_data, "total_row": new_total, "rows": rows, "layout": layout,
        "unit_first_rows": unit_first_rows, "month_cell": month_cell,
    }


def _fix_validations(ws, cols, first_data, old_end, last_data):
    targets = [cols[k] for k in ("action", "evidence") if k in cols]
    if not targets:
        return
    block = CellRange(min_col=min(targets), max_col=max(targets), min_row=first_data, max_row=max(old_end, last_data))
    best, best_area = None, -1
    for dv in list(ws.data_validations.dataValidation):
        keep = []
        for rng in list(dv.sqref.ranges):
            cr = CellRange(rng.coord)
            if cr.isdisjoint(block):
                keep.append(cr.coord)
            else:
                inter = cr.intersection(block)
                area = inter.size["rows"] * inter.size["columns"]
                # prefer a list validation whose source cells actually contain values
                src_ok = _dv_source_has_values(ws, dv)
                if src_ok and area > best_area:
                    best, best_area = dv, area
        if len(keep) != len(dv.sqref.ranges):
            if keep:
                dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(" ".join(keep))
            else:
                ws.data_validations.dataValidation.remove(dv)
    if best is None:
        return
    new = DataValidation(type=best.type, formula1=best.formula1, formula2=best.formula2,
                         allow_blank=best.allow_blank, showErrorMessage=best.showErrorMessage,
                         showInputMessage=best.showInputMessage, error=best.error,
                         errorTitle=best.errorTitle, prompt=best.prompt, promptTitle=best.promptTitle)
    new.add("%s%d:%s%d" % (col_letter(min(targets)), first_data, col_letter(max(targets)), last_data))
    ws.add_data_validation(new)


def _dv_source_has_values(ws, dv):
    f = (dv.formula1 or "").replace("$", "")
    if f.startswith('"'):
        return True
    m = re.match(r"^([A-Z]+)(\d+):([A-Z]+)(\d+)$", f)
    if not m:
        return False
    try:
        for row in ws[f]:
            for c in row:
                if c.value not in (None, ""):
                    return True
    except Exception:
        return False
    return False


def _set_month(ws, header_row, month_value, log):
    for r in range(1, header_row):
        for c in range(1, ws.max_column + 1):
            if norm_text(ws.cell(row=r, column=c).value).startswith("month"):
                for cc in range(c + 1, ws.max_column + 1):
                    cell = ws.cell(row=r, column=cc)
                    if type(cell).__name__ == "MergedCell":
                        continue
                    cell.value = month_value
                    log("MONTH\\YEAR set to %s" % month_value.strftime("%B %Y"))
                    return cell.coordinate
    log("WARNING: MONTH\\YEAR cell not found in template - not changed.")
    return None
